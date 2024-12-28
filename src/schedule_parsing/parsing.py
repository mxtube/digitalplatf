"""
Created on Thu 7.03.24
@author: Kirill Kuznetsov
"""
import openpyxl
from loguru import logger
from django.core.exceptions import ObjectDoesNotExist

from college.models import Department, CustomPerson, Auditory
from educationpart.models import Studygroup, Discipline
from schedule.models import Schedule, GroupStream, Couple

from .event import Event
from .schedule_table import ScheduleTable


class ParsingFile:

    def __init__(self, filepath: str, department: str, start_row: int = 2):
        self.department = Department.objects.get(name=department)
        self.file = openpyxl.open(filepath, read_only=True)
        self.iter_col = 2
        self.start_row = start_row
        self.error_box = []
        self.schedule = ScheduleTable()

    def start(self, start_date, end_date, weekday):
        """ Функция запуска парсинга данных из файла. Функция запускает 3 операции. """
        logger.info('Запущен парсинг расписания')
        if self._excel_reader():
            if self.__check_data_db(self.department):
                groups: list = Schedule.get_date_by_range(start_date, end_date, weekday)
                self.__save_change_gb(self.department, groups)
        return self.error_box

    def _excel_reader(self):
        """ Функция чтения данных из файла. """
        logger.info('Чтение данных из файла')
        sheet = self.file.active
        while self.iter_col < sheet.max_column:
            group_cell = sheet.cell(row=1, column=self.iter_col).value
            for iter_row in range(self.start_row, sheet.max_row + 1):
                iter_cell = sheet.cell(row=iter_row, column=self.iter_col)

                if iter_cell.value is None:
                    continue

                try:
                    couple_cell = str(sheet.cell(row=iter_row, column=1).value)
                    discipline_cell = str(sheet.cell(row=iter_row, column=self.iter_col).value.split('/')[0])
                    teacher_cell = str(sheet.cell(row=iter_row, column=self.iter_col).value.split('/')[1])
                    auditory_cell = str(sheet.cell(row=iter_row, column=self.iter_col + 1).value)
                    event = Event(
                        group=group_cell,
                        couple=couple_cell,
                        discipline=discipline_cell,
                        teacher=teacher_cell,
                        auditory=auditory_cell
                    )
                    self.schedule.add_event_to_schedule_table(event=event)
                except Exception as e:
                    logger.error(f'{e}, {iter_cell.coordinate}, {iter_cell.value}')
                    self.error_box.append(e)

            self.iter_col += 2

        count_err = len(self.error_box)
        logger.info(f'Найдено {count_err} ошибок при чтении файла')
        return count_err == 0

    def __check_data_db(self, department):
        """ Функция проверки наличия информации в базе данных """
        logger.info('Проверка расписания в базе данных')
        for group, couples in self.schedule.schedule.items():
            for couple, events in couples.items():
                for event in events:
                    try:
                        stdgrp = Studygroup.objects.get(department=department, name=group)
                        gs = GroupStream.objects.get(group=stdgrp)
                        couple = Couple.objects.get(number=couple, stream=gs.stream)

                        Discipline.objects.get(name=event.get('discipline'))
                        Auditory.objects.get(number=event.get('auditory'), department=department)
                        CustomPerson.objects.get(
                            last_name=event.get('teacher').split()[0],
                            first_name=event.get('teacher').split()[1],
                            middle_name=event.get('teacher').split()[2],
                            is_teacher=True
                        )
                    except ObjectDoesNotExist as e:
                        self.error_box.append(f'{e} {event}')
                    except Exception as e:
                        self.error_box.append(e)
        return len(self.error_box) == 0

    def __save_change_gb(self, department: Department = None, date_range: list = None):
        logger.info('Добавление расписания в базу данных')
        model = Schedule
        for date in date_range:
            logger.info('Добавление расписания на %s' % date)
            for group, couples in self.schedule.schedule.items():
                for couple, events in couples.items():
                    for event in events:
                        stdgrp = Studygroup.objects.get(department=department, name=group)
                        gs = GroupStream.objects.get(group=stdgrp)
                        grp_couple = Couple.objects.get(number=couple, stream=gs.stream)
                        discipline = Discipline.objects.get(name=event.get('discipline'))
                        auditory = Auditory.objects.get(number=event.get('auditory'), department=department)
                        teacher = CustomPerson.objects.get(
                            last_name=event.get('teacher').split()[0],
                            first_name=event.get('teacher').split()[1],
                            middle_name=event.get('teacher').split()[2],
                            is_teacher=True
                        )
                        e = model(
                            date=date,
                            couple=grp_couple,
                            group=stdgrp,
                            auditory=auditory,
                            discipline=discipline,
                            teacher=teacher
                        )
                        e.save()
                        logger.info(e)
        else:
            return True
